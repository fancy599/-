from unittest.mock import Mock, patch

from app.services.parser import ALLOWED_SUFFIXES, ParseError, extract_text_from_file


def test_allowed_suffixes_include_doc():
    assert ".doc" in ALLOWED_SUFFIXES


def test_extract_doc_calls_handler(tmp_path):
    f = tmp_path / "sample.doc"
    f.write_bytes(b"fake")
    with patch("app.services.parser._extract_doc", return_value="第十二条 测试内容足够长。"):
        text = extract_text_from_file(str(f))
    assert "第十二条" in text


def test_docx_table_keeps_row_column_relationship(tmp_path):
    from docx import Document

    f = tmp_path / "authority.docx"
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "采购金额"
    table.cell(0, 1).text = "审批主体"
    table.cell(1, 0).text = "1000万-5000万"
    table.cell(1, 1).text = "董事长"
    table.cell(2, 0).text = "5000万以上"
    table.cell(2, 1).text = "党委会前置研究、董事会审批"
    doc.save(f)

    text = extract_text_from_file(str(f))
    assert "[TABLE 1]" in text
    assert "1000万-5000万 | 董事长" in text
    assert "5000万以上 | 党委会前置研究、董事会审批" in text


def test_extract_pdf_fallback_to_ocr(tmp_path):
    f = tmp_path / "scan.pdf"
    f.write_bytes(b"fake")
    fake_doc = Mock()
    fake_doc.page_count = 1
    fake_doc.close = Mock()
    with patch("app.services.parser.fitz.open", return_value=fake_doc), patch(
        "app.services.parser._extract_pdf_native_text", return_value=[]
    ), patch("app.services.parser._extract_pdf_ocr_text", return_value=["扫描件OCR文本内容"]):
        text = extract_text_from_file(str(f))
    assert "OCR文本" in text


def test_extract_pdf_no_text_raises_helpful_error(tmp_path):
    f = tmp_path / "empty.pdf"
    f.write_bytes(b"fake")
    fake_doc = Mock()
    fake_doc.page_count = 1
    fake_doc.close = Mock()
    with patch("app.services.parser.fitz.open", return_value=fake_doc), patch(
        "app.services.parser._extract_pdf_native_text", return_value=[]
    ), patch("app.services.parser._extract_pdf_ocr_text", return_value=[]):
        try:
            extract_text_from_file(str(f))
            assert False, "expected ParseError"
        except ParseError as e:
            assert "扫描件" in str(e)
